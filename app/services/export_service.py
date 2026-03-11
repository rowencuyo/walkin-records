"""
Export service — populates the WalkIn_Records_Export.xlsx template with record data.
"""
import shutil
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from app.database import get_connection
from app.models import WalkInRecord
from app.utils.logger import get_logger

logger = get_logger(__name__)

# Path to the bundled template (project root)
_PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
TEMPLATE_PATH = _PROJECT_ROOT / "WalkIn_Records_Export.xlsx"

# Data starts on row 6 in the template
_DATA_START_ROW = 6


class ExportService:
    """Handles exporting records into the Excel template."""

    def get_records_by_date_range(
        self,
        start_date: str,
        end_date: str,
        include_inactive: bool = False,
    ) -> list[WalkInRecord]:
        """
        Fetch records whose created_at falls within [start_date, end_date].
        Dates should be in YYYY-MM-DD format.
        """
        conn = get_connection()
        conditions = []
        params = []

        if not include_inactive:
            conditions.append("is_active = 1")

        # Filter by created_at date range
        conditions.append("DATE(created_at) >= ?")
        params.append(start_date)
        conditions.append("DATE(created_at) <= ?")
        params.append(end_date)

        where = " AND ".join(conditions)
        sql = f"SELECT * FROM walkin_records WHERE {where} ORDER BY created_at ASC"
        rows = conn.execute(sql, params).fetchall()
        return [WalkInRecord.from_row(dict(r)) for r in rows]

    def export_to_excel(
        self,
        records: list[WalkInRecord],
        output_path: str,
    ) -> int:
        """
        Copy the template to output_path and populate it with the given records.
        Returns the number of records written.
        """
        if not TEMPLATE_PATH.exists():
            raise FileNotFoundError(
                f"Export template not found at {TEMPLATE_PATH}"
            )

        # Copy template to destination so we never modify the original
        shutil.copy2(str(TEMPLATE_PATH), output_path)

        wb = load_workbook(output_path)
        ws = wb.active

        for idx, rec in enumerate(records, start=1):
            row = _DATA_START_ROW + idx - 1
            ws.cell(row=row, column=1, value=idx)                          # NO
            ws.cell(row=row, column=2, value=rec.last_name)                # LASTNAME
            ws.cell(row=row, column=3, value=rec.first_name)               # FIRST NAME
            ws.cell(row=row, column=4, value=rec.middle_name)              # MIDDLE NAME
            ws.cell(row=row, column=5, value=rec.suffix_name)              # SUFFIX NAME
            ws.cell(row=row, column=6, value=rec.sex)                      # SEX
            ws.cell(row=row, column=7, value=rec.date_of_birth)            # DATE OF BIRTH
            ws.cell(row=row, column=8, value=rec.passport_number)          # PASSPORT NUMBER
            ws.cell(row=row, column=9, value=rec.country_of_citizenship)   # COUNTRY OF CITIZENSHIP
            ws.cell(row=row, column=10, value=rec.street)                  # STREET
            ws.cell(row=row, column=11, value=rec.barangay)                # BRGY
            ws.cell(row=row, column=12, value=rec.city_municipality)       # CITY/MUNICIPALITY
            ws.cell(row=row, column=13, value=rec.province)                # PROVINCE
            ws.cell(row=row, column=14, value=rec.region)                  # REGION
            ws.cell(row=row, column=15, value=rec.date_of_arrival)         # DATE OF ACCEPTANCE
            ws.cell(row=row, column=16, value=rec.date_start_education)    # DATE OF START OF CLASSES
            ws.cell(row=row, column=17, value=rec.educational_level)       # EDUCATIONAL LEVEL
            ws.cell(row=row, column=18, value=rec.course_program)          # COURSE / PROGRAM
            ws.cell(row=row, column=19, value=rec.year_level)              # YEAR LEVEL
            ws.cell(row=row, column=20, value=rec.semester)                # SEMESTER / TRIMESTER
            ws.cell(row=row, column=21, value=rec.visa_category)           # VISA CATEGORY
            ws.cell(row=row, column=22, value=rec.visa_grant_date)         # VISA GRANT DATE
            ws.cell(row=row, column=23, value=rec.visa_validity_date)      # VISA VALIDITY DATE
            ws.cell(row=row, column=24, value=rec.visa_status)             # STATUS
            ws.cell(row=row, column=25, value=rec.remarks)                 # REMARKS

        wb.save(output_path)
        wb.close()
        logger.info("Exported %d records to %s", len(records), output_path)
        return len(records)
